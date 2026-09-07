from __future__ import annotations

import io
import json
import re
from typing import Iterable


def _normalize_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _parse_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - deployment guard
        raise RuntimeError("pypdf is required for PDF ingestion") from exc

    reader = PdfReader(io.BytesIO(content))
    pages: list[str] = []

    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text)

    return "\n\n".join(pages)


def _table_rows(table) -> Iterable[str]:
    for row in table.rows:
        values = [cell.text.strip() for cell in row.cells]
        if any(values):
            yield " | ".join(values)


def _parse_docx(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - deployment guard
        raise RuntimeError("python-docx is required for DOCX ingestion") from exc

    document = Document(io.BytesIO(content))
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value:
            blocks.append(value)

    for table in document.tables:
        blocks.extend(_table_rows(table))

    return "\n\n".join(blocks)


def parse_document_bytes(
    *,
    extension: str,
    content: bytes,
) -> str:
    extension = extension.strip().lower().lstrip(".")

    if not content:
        raise ValueError("knowledge file is empty")

    if extension in {"txt", "md", "markdown", "csv"}:
        text = _decode_text(content)
    elif extension == "json":
        decoded = _decode_text(content)
        try:
            text = json.dumps(json.loads(decoded), ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            text = decoded
    elif extension == "pdf":
        text = _parse_pdf(content)
    elif extension == "docx":
        text = _parse_docx(content)
    else:
        raise ValueError(f"unsupported knowledge extension: {extension}")

    text = _normalize_text(text)
    if not text:
        raise ValueError(
            "no extractable text found; scanned PDF OCR is not enabled in P1"
        )

    return text
